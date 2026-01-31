"""Document export functionality for research results.

Supports PDF and DOCX export formats.
"""

import io
import re
from datetime import datetime
from typing import Optional

from fpdf import FPDF
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH


class ResearchPDF(FPDF):
    """Custom PDF class for research reports."""

    def __init__(self):
        super().__init__()
        self.set_auto_page_break(auto=True, margin=15)

    def header(self):
        self.set_font('Helvetica', 'B', 10)
        self.set_text_color(128, 128, 128)
        self.cell(0, 10, 'Multi-Agent Research Report', align='C', new_x='LMARGIN', new_y='NEXT')
        self.ln(5)

    def footer(self):
        self.set_y(-15)
        self.set_font('Helvetica', 'I', 8)
        self.set_text_color(128, 128, 128)
        self.cell(0, 10, f'Page {self.page_no()}', align='C')


def markdown_to_plain(text: str) -> str:
    """Convert markdown to plain text for PDF."""
    # Remove headers markers but keep text
    text = re.sub(r'^#{1,6}\s*', '', text, flags=re.MULTILINE)
    # Remove bold/italic
    text = re.sub(r'\*\*(.+?)\*\*', r'\1', text)
    text = re.sub(r'\*(.+?)\*', r'\1', text)
    text = re.sub(r'__(.+?)__', r'\1', text)
    text = re.sub(r'_(.+?)_', r'\1', text)
    # Remove links but keep text
    text = re.sub(r'\[(.+?)\]\(.+?\)', r'\1', text)
    # Remove inline code
    text = re.sub(r'`(.+?)`', r'\1', text)
    # Remove code blocks
    text = re.sub(r'```[\s\S]*?```', '', text)
    return text


def export_to_pdf(
    query: str,
    report: str,
    sources_count: Optional[int] = None,
    duration_seconds: Optional[float] = None
) -> bytes:
    """Export research results to PDF.

    Args:
        query: The research query
        report: The markdown report content
        sources_count: Number of sources used
        duration_seconds: Time taken for research

    Returns:
        PDF file as bytes
    """
    pdf = ResearchPDF()
    pdf.add_page()

    # Title
    pdf.set_font('Helvetica', 'B', 16)
    pdf.set_text_color(0, 0, 0)
    pdf.cell(0, 10, 'Research Report', new_x='LMARGIN', new_y='NEXT')
    pdf.ln(5)

    # Query
    pdf.set_font('Helvetica', 'B', 12)
    pdf.cell(0, 8, 'Query:', new_x='LMARGIN', new_y='NEXT')
    pdf.set_font('Helvetica', '', 11)
    pdf.multi_cell(0, 6, query)
    pdf.ln(5)

    # Metadata
    pdf.set_font('Helvetica', 'I', 10)
    pdf.set_text_color(100, 100, 100)
    meta_parts = [f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}"]
    if sources_count:
        meta_parts.append(f"Sources: {sources_count}")
    if duration_seconds:
        meta_parts.append(f"Duration: {duration_seconds:.1f}s")
    pdf.cell(0, 6, ' | '.join(meta_parts), new_x='LMARGIN', new_y='NEXT')
    pdf.ln(10)

    # Divider
    pdf.set_draw_color(200, 200, 200)
    pdf.line(10, pdf.get_y(), 200, pdf.get_y())
    pdf.ln(10)

    # Report content
    pdf.set_text_color(0, 0, 0)
    plain_report = markdown_to_plain(report)

    # Process content by sections
    lines = plain_report.split('\n')
    for line in lines:
        line = line.strip()
        if not line:
            pdf.ln(3)
            continue

        # Check if it looks like a heading (all caps or ends with :)
        if line.isupper() or (len(line) < 80 and line.endswith(':')):
            pdf.set_font('Helvetica', 'B', 12)
            pdf.ln(5)
            pdf.multi_cell(0, 6, line)
            pdf.set_font('Helvetica', '', 11)
        elif line.startswith('- ') or line.startswith('* '):
            # Bullet point
            pdf.set_x(15)
            pdf.multi_cell(0, 6, f"  {line}")
        else:
            pdf.multi_cell(0, 6, line)

    # Output to bytes
    return bytes(pdf.output())


def export_to_docx(
    query: str,
    report: str,
    sources_count: Optional[int] = None,
    duration_seconds: Optional[float] = None
) -> bytes:
    """Export research results to DOCX.

    Args:
        query: The research query
        report: The markdown report content
        sources_count: Number of sources used
        duration_seconds: Time taken for research

    Returns:
        DOCX file as bytes
    """
    doc = Document()

    # Title
    title = doc.add_heading('Research Report', level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Query section
    doc.add_heading('Query', level=1)
    query_para = doc.add_paragraph(query)
    query_para.italic = True

    # Metadata
    meta_para = doc.add_paragraph()
    meta_para.add_run(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}").italic = True
    if sources_count:
        meta_para.add_run(f" | Sources: {sources_count}").italic = True
    if duration_seconds:
        meta_para.add_run(f" | Duration: {duration_seconds:.1f}s").italic = True

    doc.add_paragraph()  # Spacer

    # Report content
    doc.add_heading('Findings', level=1)

    # Parse markdown and add to document
    lines = report.split('\n')
    current_para = None

    for line in lines:
        stripped = line.strip()

        if not stripped:
            current_para = None
            continue

        # Headers
        if stripped.startswith('# '):
            doc.add_heading(stripped[2:], level=1)
            current_para = None
        elif stripped.startswith('## '):
            doc.add_heading(stripped[3:], level=2)
            current_para = None
        elif stripped.startswith('### '):
            doc.add_heading(stripped[4:], level=3)
            current_para = None
        elif stripped.startswith('#### '):
            doc.add_heading(stripped[5:], level=4)
            current_para = None
        # Bullet points
        elif stripped.startswith('- ') or stripped.startswith('* '):
            doc.add_paragraph(stripped[2:], style='List Bullet')
            current_para = None
        # Numbered lists
        elif re.match(r'^\d+\.\s', stripped):
            text = re.sub(r'^\d+\.\s', '', stripped)
            doc.add_paragraph(text, style='List Number')
            current_para = None
        # Regular text
        else:
            # Handle bold and italic
            clean_text = stripped
            clean_text = re.sub(r'\*\*(.+?)\*\*', r'\1', clean_text)
            clean_text = re.sub(r'\*(.+?)\*', r'\1', clean_text)

            if current_para is None:
                current_para = doc.add_paragraph(clean_text)
            else:
                current_para.add_run(' ' + clean_text)

    # Save to bytes
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()
